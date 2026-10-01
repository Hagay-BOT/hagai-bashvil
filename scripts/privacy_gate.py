# -*- coding: utf-8 -*-
"""Privacy gate. Fails (exit 1) when anything private is about to leave the machine.

  python scripts/privacy_gate.py            -> scans files staged in git
  python scripts/privacy_gate.py dist       -> scans a build folder

Checks:
  1. any phone-like sequence (Israeli mobile / landline formats)
  2. when data-private/plan.xlsx exists: any phone number or any cell text taken from the
     private columns of the plan (lodging, water, food, notes, companions, warnings) and from
     the private tabs. The trail line itself is public OSM data and is allowed.
"""
import re
import subprocess
import sys
from pathlib import Path

PHONE = re.compile(r"(?<![\d.])0(?:5\d|[2-4689]|7\d)[-\s]?\d{3}[-\s]?\d{4}(?!\d)")
PRIVATE_COLS = (8, 9, 10, 11, 12, 13)          # lodging, water, food, notes, companions, warnings
PRIVATE_TABS = ("הטמנות מים", "מה לא אומת")
TEXT_EXT = {".html", ".js", ".mjs", ".ts", ".css", ".json", ".md", ".txt", ".sql", ".py", ".yml", ".yaml", ".svg", ".toml", ".map"}
SKIP = {"scripts/privacy_gate.py", "package-lock.json"}


def private_needles():
    src = Path("data-private/plan.xlsx")
    if not src.exists():
        return set()
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=True)
    needles = set()

    def add(cell):
        s = str(cell or "").strip()
        for m in PHONE.finditer(s):
            needles.add(re.sub(r"\D", "", m.group()))
        # sentence-sized fragments are distinctive enough to match on
        for part in re.split(r"[.;·\n]", s):
            part = part.strip()
            if len(part) >= 18:
                needles.add(part)

    ws = wb["תוכנית המסע"]
    for row in ws.iter_rows(min_row=5, values_only=True):
        for c in PRIVATE_COLS:
            if c < len(row):
                add(row[c])
    for name in PRIVATE_TABS:
        if name in wb.sheetnames:
            for row in wb[name].iter_rows(values_only=True):
                for cell in row:
                    add(cell)
    return needles


def targets(arg):
    if arg:
        return [p for p in Path(arg).rglob("*") if p.is_file()]
    out = subprocess.run(["git", "diff", "--cached", "--name-only", "--diff-filter=ACM", "-z"], capture_output=True).stdout.decode("utf-8", "replace")
    return [Path(p) for p in out.split("\0") if p]


def main():
    needles = private_needles()
    bad = []
    for p in targets(sys.argv[1] if len(sys.argv) > 1 else None):
        rel = p.as_posix()
        if rel in SKIP or p.suffix.lower() not in TEXT_EXT or not p.exists():
            continue
        text = p.read_text(encoding="utf-8", errors="replace")
        for m in PHONE.finditer(text):
            bad.append(f"{rel}: phone-like sequence «{m.group()}»")
        digits = re.sub(r"[-\s]", "", text)
        for n in needles:
            if n.isdigit():
                if n in digits:
                    bad.append(f"{rel}: phone number from the private sheet")
            elif n in text:
                bad.append(f"{rel}: private sheet text «{n[:24]}…»")
    if bad:
        print("PRIVACY GATE FAILED")
        for b in bad[:40]:
            print("  " + b)
        return 1
    print(f"privacy gate ok ({len(needles)} private fragments checked)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
